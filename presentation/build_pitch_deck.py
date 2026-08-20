from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR

NVIDIA_GREEN = RGBColor(0x76, 0xB9, 0x00)
DARK_NAVY = RGBColor(0x0B, 0x17, 0x26)
SLATE = RGBColor(0x3A, 0x47, 0x59)
LIGHT_GRAY = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x1F, 0x4E, 0x8C)
RED = RGBColor(0xB4, 0x23, 0x2E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_slide():
    return prs.slides.add_slide(BLANK)

def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def add_title(slide, text, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.28), Inches(1.15), Inches(0.07))
    bar.fill.solid(); bar.fill.fore_color.rgb = NVIDIA_GREEN; bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.42), Inches(12.3), Inches(0.8))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = DARK_NAVY
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.5))
        tf2 = box2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; r2 = p2.add_run(); r2.text = subtitle
        r2.font.size = Pt(15); r2.font.italic = True; r2.font.color.rgb = SLATE

def text_block(slide, left, top, width, height, lines, size=15, color=DARK_NAVY, bold_head=None, bullet=True, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = ("- " + line) if bullet else line
        r.font.size = Pt(size); r.font.color.rgb = color
    return box

def stat_card(slide, left, top, width, height, number, label, accent=NVIDIA_GREEN):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT_GRAY
    card.line.color.rgb = accent; card.line.width = Pt(1.5)
    card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = number
    r1.font.size = Pt(30); r1.font.bold = True; r1.font.color.rgb = accent
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label
    r2.font.size = Pt(12); r2.font.color.rgb = DARK_NAVY

def flow_box(slide, left, top, width, height, text, fill=BLUE, font_color=WHITE, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = font_color
    return shp

def down_arrow(slide, x_center, top, length=Inches(0.18)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x_center, top, x_center, top + length)
    conn.line.color.rgb = SLATE
    conn.line.width = Pt(2.25)
    line = conn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = line.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    line.append(tail)

# ---------------------------------------------------------------------------
# Slide 1 - The challenge
# ---------------------------------------------------------------------------
s1 = add_slide(); set_bg(s1)
add_title(s1, "The Challenge: Recalls Are Discovered Too Late",
          "Vehicle Defect Pattern Agent - RecallZero")
text_block(s1, Inches(0.6), Inches(1.85), Inches(6.3), Inches(4.6), [
    "Consumer safety complaints arrive by the thousands per vehicle, in free text, from unrelated owners.",
    "Quality and safety engineers manually sift this volume with keyword search and periodic review.",
    "By the time a pattern is obvious enough to notice manually, a formal recall is often already underway - the warning arrives after the fact.",
    "There is no consistent, repeatable, auditable way to flag an emerging defect before it becomes a recall.",
], size=16)
text_block(s1, Inches(7.2), Inches(1.85), Inches(5.5), Inches(4.6), [
    "Who feels this problem",
], size=17, bold_head=True, bullet=False, color=NVIDIA_GREEN)
card1 = slide_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.35), Inches(5.5), Inches(3.9))
card1.fill.solid(); card1.fill.fore_color.rgb = LIGHT_GRAY; card1.line.color.rgb = NVIDIA_GREEN
card1.shadow.inherit = False
tf = card1.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.2)
items = [
    "OEM quality and safety engineering teams triaging fleet-wide complaint volume",
    "Regulatory and compliance reviewers who need a defensible evidence trail",
    "Investigation leads deciding whether an issue is worth opening a formal review",
]
for i, item in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(14)
    r = p.add_run(); r.text = "- " + item
    r.font.size = Pt(16); r.font.color.rgb = DARK_NAVY

# ---------------------------------------------------------------------------
# Slide 2 - Landscape and differentiation
# ---------------------------------------------------------------------------
s2 = add_slide(); set_bg(s2)
add_title(s2, "Landscape and Differentiation", "What exists today, and why RecallZero is different")

rows = [
    ("Today", "RecallZero"),
    ("Manual keyword search across NHTSA complaints", "Automatic NHTSA ingestion, clustered by component and root failure mechanism"),
    ("Generic AI summarizers can invent counts, trends, or causation", "AI (NVIDIA NIM) only interprets language per complaint; every count, trend, and score is deterministic code"),
    ("Severity and priority are an opaque single score", "Explainable 5-factor weighted risk score - every contribution is printed"),
    ("Recall status is checked manually, after the fact", "Recall cross-reference actively changes the score - already-covered issues stop alerting"),
    ("Early-warning claims are asserted, not tested", "Leakage-safe historical replay proved a 15-day early alert on a real recall, then the detector was frozen and stress-tested on 20 held-out cases"),
]
table_shape = s2.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.3))
table = table_shape.table
table.columns[0].width = Inches(5.6)
table.columns[1].width = Inches(6.5)
for r_idx, (left_text, right_text) in enumerate(rows):
    for c_idx, txt in enumerate((left_text, right_text)):
        cell = table.cell(r_idx, c_idx)
        cell.text = txt
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
        para = cell.text_frame.paragraphs[0]
        run = para.runs[0]
        if r_idx == 0:
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = WHITE
            cell.fill.solid(); cell.fill.fore_color.rgb = DARK_NAVY
        else:
            run.font.size = Pt(13.5)
            run.font.color.rgb = DARK_NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if c_idx == 0 else RGBColor(0xEE, 0xF6, 0xDC)

# ---------------------------------------------------------------------------
# Slide 3 - Customer benefit and ROI
# ---------------------------------------------------------------------------
s3 = add_slide(); set_bg(s3)
add_title(s3, "Customer Benefit and ROI", "Value delivered, measured under leakage-safe historical replay")

stats = [
    ("15 days", "Earlier alert than the official recall date, measured on a real historical campaign (Ford Mustang Mach-E, 22V412000)"),
    ("~11%", "Unconfirmed-alert rate on clean control vehicles with no matching recall - low false-alarm burden"),
    ("100%", "Of alerts trace to real NHTSA complaint IDs - fully auditable, not a black-box score"),
    ("1,000s -> few", "Raw complaints per vehicle condensed into a short, ranked list of engineering signals"),
]
card_w = Inches(2.9); card_h = Inches(2.0); gap = Inches(0.25)
start_x = Inches(0.6); y = Inches(1.9)
for i, (num, label) in enumerate(stats):
    x = start_x + i * (card_w + gap)
    stat_card(s3, x, y, card_w, card_h, num, label, accent=NVIDIA_GREEN if i % 2 == 0 else BLUE)

text_block(s3, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.6), [
    "Faster triage: engineers review a short, ranked brief instead of reading thousands of raw complaints.",
    "Lower compliance risk: every alert is traceable to source evidence and every score shows its calculation - defensible in an audit.",
    "Earlier intervention: a validated earlier signal gives quality teams more time to investigate before a defect pattern grows to recall scale.",
    "Honest confidence: the detector was frozen and validated against 10 real recalls and 10 negative controls before any accuracy claim was made.",
], size=15.5)

# ---------------------------------------------------------------------------
# Slide 4 - Architecture diagram
# ---------------------------------------------------------------------------
s4 = add_slide(); set_bg(s4)
add_title(s4, "Architecture", "AI interprets complaint language - deterministic code decides risk, timing, and evidence")

flow_x = Inches(0.6); flow_w = Inches(6.6)
flow_h = Inches(0.68)
ys = [Inches(1.85), Inches(2.68), Inches(3.51), Inches(4.34), Inches(5.17), Inches(6.00)]
labels = [
    ("NHTSA Complaints + Recalls (public API)", SLATE),
    ("NVIDIA NIM - Structured Failure-Signature Extraction", NVIDIA_GREEN),
    ("NVIDIA Embeddings + Deterministic Component/Taxonomy Clustering", NVIDIA_GREEN),
    ("Deterministic Risk Engine - Severity / Trend / Persistence / Evidence / Recall-Gap", BLUE),
    ("Recall Cross-Reference (visible recalls only at cutoff)", BLUE),
    ("Engineering Brief + Alert", DARK_NAVY),
]
for (text, color), top in zip(labels, ys):
    flow_box(s4, flow_x, top, flow_w, flow_h, text, fill=color)
for top in ys[:-1]:
    down_arrow(s4, flow_x + flow_w / 2, top + flow_h)

agent_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.55), Inches(1.85), Inches(5.2), Inches(2.6))
agent_box.fill.solid(); agent_box.fill.fore_color.rgb = WHITE
agent_box.line.color.rgb = NVIDIA_GREEN; agent_box.line.width = Pt(2)
agent_box.line.dash_style = None
tf = agent_box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
p0 = tf.paragraphs[0]; r0 = p0.add_run(); r0.text = "NeMo Agent Toolkit Agent"
r0.font.bold = True; r0.font.size = Pt(15); r0.font.color.rgb = NVIDIA_GREEN
for item in ["Calls RecallZero tools on a natural-language request",
             "Narrates already-computed results",
             "Never invents counts, trends, or recall status"]:
    p = tf.add_paragraph(); p.space_before = Pt(8)
    r = p.add_run(); r.text = "- " + item
    r.font.size = Pt(13); r.font.color.rgb = DARK_NAVY

validation_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.55), Inches(4.65), Inches(5.2), Inches(2.0))
validation_box.fill.solid(); validation_box.fill.fore_color.rgb = LIGHT_GRAY
validation_box.line.color.rgb = BLUE; validation_box.line.width = Pt(2)
tf = validation_box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
p0 = tf.paragraphs[0]; r0 = p0.add_run(); r0.text = "Recall Time Machine"
r0.font.bold = True; r0.font.size = Pt(15); r0.font.color.rgb = BLUE
for item in ["Leakage-safe weekly historical replay",
             "Detector Freeze v1 + 20-case validation cohort",
             "Pre-registered acceptance criteria"]:
    p = tf.add_paragraph(); p.space_before = Pt(8)
    r = p.add_run(); r.text = "- " + item
    r.font.size = Pt(13); r.font.color.rgb = DARK_NAVY

bottom = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.85), Inches(12.1), Inches(0.5))
bottom.fill.solid(); bottom.fill.fore_color.rgb = DARK_NAVY; bottom.line.fill.background()
tf = bottom.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Trust boundary: AI understands language. Deterministic code counts evidence, computes trend, and decides risk."
r.font.size = Pt(13); r.font.color.rgb = WHITE; r.font.bold = True

prs.save("presentation/RecallZero_Pitch.pptx")
print("SAVED")